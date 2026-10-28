"""OSS text extraction — pipeline layer 1 (spec §7.3, research 02).

Free and local: pypdf as a fast pre-pass (encryption/corruption check) into
pdfplumber (layout-aware) for PDFs; python-docx for DOCX with a raw
``word/document.xml`` strip fallback (text boxes); TXT is a decode. Scanned
pages (no text layer) are detected by text density. The document preparer in
``documents.py`` handles OCR, vision and gated legacy-format conversion.

Every unusable document raises :class:`ExtractionError` with a stable machine
reason; the pipeline maps that to an explicit terminal status. Nothing is
dropped silently.
"""

import io
import re
import unicodedata
import zipfile

from app.config import settings

# A PDF page with fewer visible characters than this is "scanned" (research 02).
_SCANNED_PAGE_CHARS = 50


class ExtractionError(Exception):
    """The text layer is unusable — reason is a stable machine string."""

    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


def sniff_kind(data: bytes, filename: str) -> str:
    """Classify PDF, DOCX, text, raster images, OLE DOC and RTF by content.

    Raises :class:`ExtractionError` (``unsupported_format``) for anything the
    document preparer cannot parse — the GUI nudges toward PDF/DOCX/TXT.
    """
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "doc"
    if data.lstrip().startswith(b"{\\rtf"):
        return "rtf"
    if data.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"II*\x00", b"MM\x00*")):
        return "image"
    if data.startswith(b"RIFF") and data[8:12] == b"WEBP":
        return "image"
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ExtractionError("document_too_large")
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            raise ExtractionError("unsupported_format") from None
        raise ExtractionError("unsupported_format")
    try:
        data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ExtractionError("unsupported_format") from None
    if b"\x00" in data or any(byte < 9 for byte in data):
        raise ExtractionError("unsupported_format")
    return "txt"


class ExtractedText:
    __slots__ = ("text", "truncated", "raw_text")

    def __init__(self, text: str, truncated: bool = False, raw_text: str | None = None):
        self.text = text
        self.truncated = truncated
        self.raw_text = text if raw_text is None else raw_text


def extract_text(data: bytes, filename: str) -> ExtractedText:
    """Extract resume text from original bytes; values stay verbatim."""
    kind = sniff_kind(data, filename)
    if kind == "pdf":
        text = _extract_pdf(data)
    elif kind == "docx":
        text = _extract_docx(data)
    elif kind == "txt":
        text = _extract_txt(data)
    else:
        raise ExtractionError("unsupported_format")
    text = _normalize(text)
    if not text.strip():
        raise ExtractionError("empty_text")
    if len(text) > settings.resume_text_cap_chars:
        return ExtractedText(text[: settings.resume_text_cap_chars], truncated=True, raw_text=text)
    return ExtractedText(text)


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------


def _extract_pdf(data: bytes) -> str:
    # pypdf pre-pass: cheap encryption/corruption gate (research 02 §5).
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractionError("encrypted")
            except Exception:  # noqa: BLE001 — any decrypt failure
                raise ExtractionError("encrypted") from None
        if len(reader.pages) > settings.resume_max_pages:
            raise ExtractionError("too_many_pages")
    except ExtractionError:
        raise
    except Exception:  # noqa: BLE001 — pypdf rejects it: try pdfplumber below
        pass

    # pdfplumber primary extractor — layout-aware reading order.
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
    except Exception:  # noqa: BLE001 — both layers failed to parse
        raise ExtractionError("corrupt_pdf") from None

    if not pages:
        raise ExtractionError("corrupt_pdf")
    if any(len(page.strip()) < _SCANNED_PAGE_CHARS for page in pages):
        # The document preparer tries OCR, then vision for these pages.
        raise ExtractionError("scanned_or_empty")
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    parts: list[str] = []
    try:
        document = docx.Document(io.BytesIO(data))
        parts.extend(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    except Exception:  # noqa: BLE001 — fall through to the raw-XML strip
        pass

    if not any(part.strip() for part in parts):
        # Text boxes (and other structures python-docx misses) still carry
        # <w:t> nodes in word/document.xml — strip them raw (research 02 §6.5).
        parts.append(_strip_docx_xml(data))

    text = "\n".join(part for part in parts if part.strip())
    if not text.strip():
        raise ExtractionError("docx_empty")
    return text


def _strip_docx_xml(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", "replace")
    except zipfile.BadZipFile, KeyError:
        return ""
    return "\n".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8-sig")


def _normalize(text: str) -> str:
    """NFKC + whitespace cleanup (research 02 §5). ZWNJ survives NFKC, so
    Persian text keeps its zero-width non-joiners verbatim."""
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[ \t ]+", " ", text, flags=re.MULTILINE)
    text = re.sub(r"^ +| +$", "", text, flags=re.MULTILINE)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
